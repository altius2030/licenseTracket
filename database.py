"""
SQLite database module for License Tracker.
Handles all database operations: create tables, CRUD, import/export.
"""
import sqlite3
import os
import sys
from datetime import datetime, date, timedelta
from typing import List, Dict, Any, Optional, Tuple
from contextlib import contextmanager
import json


def get_app_dir() -> str:
    """Return the directory where the app/exe is located, for portable DB storage."""
    if getattr(sys, 'frozen', False):
        # Running as compiled exe
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


class Database:
    def __init__(self, db_path: str = None):
        if db_path is None:
            db_path = os.path.join(get_app_dir(), "licenses.db")
        self.db_path = db_path
        self.init_db()

    @contextmanager
    def get_connection(self):
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        try:
            yield conn
        finally:
            conn.close()

    def init_db(self):
        """Create tables if they don't exist."""
        with self.get_connection() as conn:
            cursor = conn.cursor()

            # Main licenses table - unified schema covering all sheets
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS licenses (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    category TEXT NOT NULL DEFAULT 'GBS-market',
                    sheet_name TEXT NOT NULL,
                    row_number INTEGER,
                    
                    -- Core fields (from GBS sheet)
                    install_date DATE,
                    full_name TEXT,
                    store_name TEXT,
                    phone TEXT,
                    email TEXT,
                    gbs_id TEXT,
                    activation_period TEXT,
                    id_expiry DATE,
                    payment_date DATE,
                    renewal_date DATE,
                    next_payment DATE,
                    address TEXT,
                    
                    -- Additional fields (from other sheets)
                    license_key TEXT,
                    payment_info TEXT,
                    receipt_number TEXT,
                    migrated_from TEXT,
                    comment TEXT,
                    sent_info TEXT,
                    activation_info TEXT,
                    
                    -- Status tracking
                    status TEXT DEFAULT 'active',  -- active, expiring, expired, manual
                    manual_status TEXT,           -- user override: active, expiring, expired
                    status_updated_at TIMESTAMP,
                    
                    -- Metadata
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            """)

            # Indexes for performance
            cursor.execute("CREATE INDEX IF NOT EXISTS idx_sheet_name ON licenses(sheet_name)")
            cursor.execute("CREATE INDEX IF NOT EXISTS idx_full_name ON licenses(full_name)")
            cursor.execute("CREATE INDEX IF NOT EXISTS idx_next_payment ON licenses(next_payment)")
            cursor.execute("CREATE INDEX IF NOT EXISTS idx_id_expiry ON licenses(id_expiry)")
            cursor.execute("CREATE INDEX IF NOT EXISTS idx_status ON licenses(status)")

            # Settings table
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS settings (
                    key TEXT PRIMARY KEY,
                    value TEXT
                )
            """)

            # Import history
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS import_history (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    file_path TEXT,
                    sheet_name TEXT,
                    records_imported INTEGER,
                    records_updated INTEGER,
                    records_skipped INTEGER,
                    imported_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            """)

            # Comments
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS comments (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    license_id INTEGER NOT NULL,
                    text TEXT NOT NULL,
                    created_at TIMESTAMP,
                    updated_at TIMESTAMP
                )
            """)
            cursor.execute("CREATE INDEX IF NOT EXISTS idx_comments_license ON comments(license_id)")

            conn.commit()

    # ===== CRUD Operations =====

    def get_all_licenses(self, category: Optional[str] = None, sheet_name: Optional[str] = None) -> List[Dict]:
        """Get all licenses, optionally filtered by category and sheet."""
        with self.get_connection() as conn:
            cursor = conn.cursor()
            query = "SELECT * FROM licenses WHERE 1=1"
            params = []
            if category:
                query += " AND category = ?"
                params.append(category)
            if sheet_name:
                query += " AND sheet_name = ?"
                params.append(sheet_name)
            query += " ORDER BY id"
            cursor.execute(query, params)
            return [dict(row) for row in cursor.fetchall()]

    def get_license(self, license_id: int) -> Optional[Dict]:
        """Get a single license by ID."""
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM licenses WHERE id = ?", (license_id,))
            row = cursor.fetchone()
            return dict(row) if row else None

    def get_categories(self) -> List[str]:
        """Get list of distinct categories."""
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT DISTINCT category FROM licenses ORDER BY category")
            return [row[0] for row in cursor.fetchall()]
    
    def get_sheets(self, category: Optional[str] = None) -> List[str]:
        """Get list of distinct sheet names, optionally filtered by category."""
        with self.get_connection() as conn:
            cursor = conn.cursor()
            if category:
                cursor.execute("SELECT DISTINCT sheet_name FROM licenses WHERE category = ? ORDER BY sheet_name", (category,))
            else:
                cursor.execute("SELECT DISTINCT sheet_name FROM licenses ORDER BY sheet_name")
            return [row[0] for row in cursor.fetchall()]

    def create_license(self, data: Dict[str, Any]) -> int:
        """Create a new license record."""
        data['created_at'] = datetime.now().isoformat()
        data['updated_at'] = datetime.now().isoformat()
        data['status_updated_at'] = datetime.now().isoformat()
        
        # Calculate initial status
        data['status'] = self._calculate_status(data)
        
        fields = ', '.join(data.keys())
        placeholders = ', '.join(['?' for _ in data])
        values = list(data.values())
        
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(f"INSERT INTO licenses ({fields}) VALUES ({placeholders})", values)
            conn.commit()
            return cursor.lastrowid

    def update_license(self, license_id: int, data: Dict[str, Any]) -> bool:
        """Update a license record."""
        data['updated_at'] = datetime.now().isoformat()
        
        # Recalculate status unless manual override exists
        if 'manual_status' not in data or data['manual_status'] is None:
            # Merge existing record data with new data for accurate status calculation
            existing = self.get_license(license_id)
            merged = {**existing, **data} if existing else data
            data['status'] = self._calculate_status(merged)
        else:
            data['status'] = data['manual_status']
            data['status_updated_at'] = datetime.now().isoformat()
        
        set_clause = ', '.join([f"{k} = ?" for k in data.keys()])
        values = list(data.values()) + [license_id]
        
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(f"UPDATE licenses SET {set_clause} WHERE id = ?", values)
            conn.commit()
            return cursor.rowcount > 0

    def delete_license(self, license_id: int) -> bool:
        """Delete a license record and its comments."""
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("DELETE FROM comments WHERE license_id = ?", (license_id,))
            cursor.execute("DELETE FROM licenses WHERE id = ?", (license_id,))
            conn.commit()
            return cursor.rowcount > 0

    def move_licenses(self, license_ids: List[int], category: str) -> int:
        """Move records to another category. Returns the number of moved rows."""
        ids = [int(i) for i in license_ids]
        if not ids or not category:
            return 0

        now = datetime.now().isoformat()
        total = 0
        with self.get_connection() as conn:
            cursor = conn.cursor()
            # Chunk to stay below SQLite's bound-variable limit
            for start in range(0, len(ids), 500):
                chunk = ids[start:start + 500]
                placeholders = ', '.join(['?' for _ in chunk])
                cursor.execute(
                    f"UPDATE licenses SET category = ?, updated_at = ? "
                    f"WHERE id IN ({placeholders})",
                    [category, now] + chunk
                )
                total += cursor.rowcount
            conn.commit()
        return total

    def move_sheet_records(self, sheet_name: str, from_category: str, to_category: str) -> int:
        """Move all records of one sheet from one category to another."""
        now = datetime.now().isoformat()
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(
                "UPDATE licenses SET category = ?, updated_at = ? "
                "WHERE sheet_name = ? AND category = ?",
                (to_category, now, sheet_name, from_category)
            )
            conn.commit()
            return cursor.rowcount

    # ===== Comments =====

    def get_comments(self, license_id: int) -> List[Dict]:
        """Get all comments for a license, newest first."""
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(
                "SELECT * FROM comments WHERE license_id = ? ORDER BY created_at DESC, id DESC",
                (license_id,)
            )
            return [dict(row) for row in cursor.fetchall()]

    def add_comment(self, license_id: int, text: str) -> int:
        """Add a comment with current timestamp."""
        now = datetime.now().isoformat()
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(
                "INSERT INTO comments (license_id, text, created_at, updated_at) VALUES (?, ?, ?, ?)",
                (license_id, text, now, now)
            )
            conn.commit()
            return cursor.lastrowid

    def update_comment(self, comment_id: int, text: str) -> bool:
        """Edit a comment, refreshing updated_at."""
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(
                "UPDATE comments SET text = ?, updated_at = ? WHERE id = ?",
                (text, datetime.now().isoformat(), comment_id)
            )
            conn.commit()
            return cursor.rowcount > 0

    def delete_comment(self, comment_id: int) -> bool:
        """Delete a comment."""
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("DELETE FROM comments WHERE id = ?", (comment_id,))
            conn.commit()
            return cursor.rowcount > 0

    def bulk_upsert(self, records: List[Dict], sheet_name: str, category: str = 'GBS-market') -> Tuple[int, int, int]:
        """
        Bulk insert/update records from Excel import.
        Returns: (imported, updated, skipped)
        """
        imported = 0
        updated = 0
        skipped = 0
        
        with self.get_connection() as conn:
            cursor = conn.cursor()
            
            for record in records:
                record['sheet_name'] = sheet_name
                record['category'] = category
                record['status'] = self._calculate_status(record)
                record['created_at'] = datetime.now().isoformat()
                record['updated_at'] = datetime.now().isoformat()
                record['status_updated_at'] = datetime.now().isoformat()
                
                # Check if record exists (by gbs_id, license_key, or fallback key)
                existing_id = None
                if record.get('gbs_id'):
                    cursor.execute(
                        "SELECT id FROM licenses WHERE sheet_name = ? AND gbs_id = ?",
                        (sheet_name, record['gbs_id'])
                    )
                    row = cursor.fetchone()
                    if row:
                        existing_id = row[0]
                elif record.get('license_key'):
                    cursor.execute(
                        "SELECT id FROM licenses WHERE sheet_name = ? AND license_key = ?",
                        (sheet_name, record['license_key'])
                    )
                    row = cursor.fetchone()
                    if row:
                        existing_id = row[0]
                elif record.get('full_name'):
                    # Fallback: match by full_name + store_name within same sheet
                    cursor.execute(
                        "SELECT id FROM licenses WHERE sheet_name = ? AND full_name = ? AND store_name = ?",
                        (sheet_name, record['full_name'], record.get('store_name', ''))
                    )
                    row = cursor.fetchone()
                    if row:
                        existing_id = row[0]
                
                if existing_id:
                    # Update existing
                    set_clause = ', '.join([f"{k} = ?" for k in record.keys()])
                    values = list(record.values()) + [existing_id]
                    cursor.execute(f"UPDATE licenses SET {set_clause} WHERE id = ?", values)
                    updated += 1
                else:
                    # Insert new
                    fields = ', '.join(record.keys())
                    placeholders = ', '.join(['?' for _ in record])
                    values = list(record.values())
                    cursor.execute(f"INSERT INTO licenses ({fields}) VALUES ({placeholders})", values)
                    imported += 1
            
            conn.commit()
        
        return imported, updated, skipped

    def _calculate_status(self, data: Dict) -> str:
        """Auto-calculate license status based on dates."""
        today = date.today()
        
        def evaluate(value):
            if value is None:
                return None
            if isinstance(value, str):
                try:
                    value = datetime.fromisoformat(value.replace('Z', '')).date()
                except:
                    return None
            if isinstance(value, datetime):
                value = value.date()
            if isinstance(value, date):
                d = (value - today).days
                if d < 0:
                    return 'expired'
                elif d <= 30:
                    return 'expiring'
                return 'active'
            return None
        
        # Primary source: Срок ID (id_expiry)
        result = evaluate(data.get('id_expiry'))
        if result:
            return result
        
        # Fallback: next_payment when id_expiry is missing
        result = evaluate(data.get('next_payment'))
        if result:
            return result
        
        return 'active'

    def refresh_all_statuses(self) -> int:
        """Recalculate status for all licenses (respecting manual overrides)."""
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT id, next_payment, id_expiry, manual_status FROM licenses")
            rows = cursor.fetchall()
            
            updated = 0
            for row in rows:
                if row['manual_status']:
                    continue  # Skip manual overrides
                
                data = {'next_payment': row['next_payment'], 'id_expiry': row['id_expiry']}
                new_status = self._calculate_status(data)
                
                if new_status != row['status']:
                    cursor.execute(
                        "UPDATE licenses SET status = ?, status_updated_at = ? WHERE id = ?",
                        (new_status, datetime.now().isoformat(), row['id'])
                    )
                    updated += 1
            
            conn.commit()
            return updated

    def get_expiring_soon(self, days: int = 30) -> List[Dict]:
        """Get licenses expiring within specified days."""
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                SELECT * FROM licenses 
                WHERE (next_payment IS NOT NULL AND next_payment <= date('now', ?) AND next_payment >= date('now'))
                   OR (id_expiry IS NOT NULL AND id_expiry <= date('now', ?) AND id_expiry >= date('now'))
                ORDER BY COALESCE(next_payment, id_expiry)
            """, (f'+{days} days', f'+{days} days'))
            return [dict(row) for row in cursor.fetchall()]

    def get_expired(self) -> List[Dict]:
        """Get all expired licenses."""
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                SELECT * FROM licenses 
                WHERE (next_payment IS NOT NULL AND next_payment < date('now'))
                   OR (id_expiry IS NOT NULL AND id_expiry < date('now'))
                ORDER BY COALESCE(next_payment, id_expiry)
            """)
            return [dict(row) for row in cursor.fetchall()]

    # ===== Settings =====

    def get_setting(self, key: str, default: Any = None) -> Any:
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT value FROM settings WHERE key = ?", (key,))
            row = cursor.fetchone()
            if row:
                try:
                    return json.loads(row[0])
                except:
                    return row[0]
            return default

    def set_setting(self, key: str, value: Any):
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(
                "INSERT OR REPLACE INTO settings (key, value) VALUES (?, ?)",
                (key, json.dumps(value))
            )
            conn.commit()

    # ===== Import History =====

    def log_import(self, file_path: str, sheet_name: str, imported: int, updated: int, skipped: int):
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                INSERT INTO import_history (file_path, sheet_name, records_imported, records_updated, records_skipped)
                VALUES (?, ?, ?, ?, ?)
            """, (file_path, sheet_name, imported, updated, skipped))
            conn.commit()

    def get_import_history(self) -> List[Dict]:
        with self.get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM import_history ORDER BY imported_at DESC")
            return [dict(row) for row in cursor.fetchall()]

    # ===== Export =====

    def export_to_excel(self, file_path: str, category: Optional[str] = None, sheet_name: Optional[str] = None):
        """Export licenses to Excel file with separate sheets."""
        import pandas as pd
        
        licenses = self.get_all_licenses(category, sheet_name)
        if not licenses:
            return
        
        df = pd.DataFrame(licenses)
        
        # Drop internal columns
        drop_cols = ['id', 'created_at', 'updated_at', 'status_updated_at', 'row_number', 'category']
        df = df.drop(columns=[c for c in drop_cols if c in df.columns])
        
        with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
            if sheet_name:
                df.to_excel(writer, sheet_name=sheet_name, index=False)
            else:
                # Group by sheet_name
                for sn in df['sheet_name'].unique():
                    sheet_df = df[df['sheet_name'] == sn].drop(columns=['sheet_name'])
                    sheet_df.to_excel(writer, sheet_name=sn[:31], index=False)  # Excel sheet name limit


# Column mapping for Excel import
EXCEL_COLUMN_MAP = {
    # GBS sheet
    'Дата установки': 'install_date',
    'Фамилия и имя': 'full_name',
    'Название магазина': 'store_name',
    'Телефон': 'phone',
    'Почта': 'email',
    'GBS ID': 'gbs_id',
    'Срок активации ': 'activation_period',
    'Срок  ID': 'id_expiry',
    'Дата оплаты ': 'payment_date',
    'Дата продления ': 'renewal_date',
    'Следующая оплата ': 'next_payment',
    'Адрес': 'address',
    
    # ШопУчет sheet
    'ФИО': 'full_name',
    'Магазин': 'store_name',
    'ключ': 'license_key',
    'срок когла оплатил ': 'payment_date',
    
    # шоп учет sheet
    '№': 'row_number',
    'Отправлено': 'sent_info',
    'Активация и срок': 'activation_info',
    'Оплата': 'payment_info',
    '№ квитанции': 'receipt_number',
    'Переход с другой программы': 'migrated_from',
    'Комментарий': 'comment',
    
    # АСТАНА sheet (same as шоп учет mostly)
    # Алматы sheet (same as GBS mostly)
    # отвалы sheet (similar to шоп учет with many extra cols)
}


def parse_excel_date(value) -> Optional[str]:
    """Parse various date formats from Excel to ISO string."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        # Excel serial date
        try:
            return (datetime(1899, 12, 30) + timedelta(days=value)).date().isoformat()
        except:
            return None
    if isinstance(value, str):
        # Try various formats
        for fmt in ['%d.%m.%Y', '%Y-%m-%d', '%d/%m/%Y', '%d.%m.%y', '%d %m %Y']:
            try:
                return datetime.strptime(value.strip(), fmt).date().isoformat()
            except:
                continue
        # Try parsing as datetime
        try:
            return datetime.fromisoformat(value.replace('Z', '')).date().isoformat()
        except:
            pass
    return None


def import_excel_file(db: Database, file_path: str, category: str = 'GBS-market') -> Dict[str, Tuple[int, int, int]]:
    """Import all sheets from Excel file. Returns dict of sheet_name -> (imported, updated, skipped)."""
    import openpyxl
    
    wb = openpyxl.load_workbook(file_path, data_only=True)
    results = {}
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        records = []
        
        # Get headers from first row
        headers = []
        for cell in ws[1]:
            headers.append(cell.value)
        
        # Map headers to database fields
        col_map = {}
        for idx, header in enumerate(headers):
            if header and header in EXCEL_COLUMN_MAP:
                col_map[idx] = EXCEL_COLUMN_MAP[header]
        
        # Read data rows
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
            record = {'row_number': row_idx}
            has_data = False
            
            for col_idx, field_name in col_map.items():
                cell = row[col_idx] if col_idx < len(row) else None
                value = cell.value if cell else None
                
                if value is not None:
                    has_data = True
                    # Parse dates
                    if field_name in ['install_date', 'id_expiry', 'payment_date', 'renewal_date', 'next_payment']:
                        value = parse_excel_date(value)
                    record[field_name] = value
            
            if has_data:
                records.append(record)
        
        if records:
            imported, updated, skipped = db.bulk_upsert(records, sheet_name, category)
            db.log_import(file_path, sheet_name, imported, updated, skipped)
            results[sheet_name] = (imported, updated, skipped)
    
    return results